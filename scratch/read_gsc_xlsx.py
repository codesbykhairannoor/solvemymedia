import os
import openpyxl
import pandas as pd

folder = r"d:\audiovideo\media-compressor\csvsolve"
files = [
    "solvemymedia.com-Coverage-Drilldown-2026-09-17.xlsx",
    "solvemymedia.com-Coverage-Drilldown-2026-09-17(1).xlsx",
    "solvemymedia.com-Coverage-Drilldown-2026-09-17(2).xlsx",
    "solvemymedia.com-Coverage-Drilldown-2026-09-17(3).xlsx"
]

for filename in files:
    filepath = os.path.join(folder, filename)
    print("=" * 80)
    print(f"FILE: {filename}")
    print("=" * 80)
    wb = openpyxl.load_workbook(filepath, data_only=True)
    print(f"Sheet names: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        df = pd.read_excel(filepath, sheet_name=sheet_name)
        print(f"\n--- Sheet: {sheet_name} (Shape: {df.shape}) ---")
        print("Columns:", list(df.columns))
        print("Head:")
        print(df.head(10).to_string())
        print("\nTail (if > 10 rows):")
        if len(df) > 10:
            print(df.tail(5).to_string())
        print("-" * 40)
